"""Untrusted supplementary evidence ingestion, without storage or publication.

Only the caller decides how to retain these dictionaries. Roles are provenance,
not permission to publish: policy/reference/student text remains untrusted input.
"""
from contextlib import contextmanager
import hashlib
from html.parser import HTMLParser
import http.client
import io
import ipaddress
import os
from pathlib import Path, PurePosixPath
import socket
import ssl
import stat
import tempfile
from urllib.parse import quote, urljoin, urlsplit, urlunsplit
import zipfile


MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_TEXT_CHARS = 60000
MAX_ZIP_BYTES = 32 * 1024 * 1024
MAX_ZIP_ENTRIES = 2000
MAX_ZIP_RATIO = 200
MAX_RESPONSE_BYTES = 2 * 1024 * 1024
MAX_REDIRECTS = 3
CONNECT_TIMEOUT = 5
READ_TIMEOUT = 10
ROLES = frozenset(("reference", "policy", "student"))
EXTENSIONS = frozenset((".txt", ".md", ".csv", ".html", ".pdf", ".docx", ".pptx"))


class MaterialError(ValueError):
    """A safe user-facing material ingestion error."""


def _role(role):
    if not isinstance(role, str) or role not in ROLES:
        raise MaterialError("Choose a reference, policy, or student material role.")
    return role


def _warnings(role):
    if role == "policy":
        return ["Policy material informs course design; it is not automatically published."]
    return []


def _name(value):
    if not isinstance(value, (str, os.PathLike)):
        raise MaterialError("Provide a material name.")
    name = str(value).replace("\\", "/").split("/")[-1]
    name = "".join(c for c in name if c.isprintable()).strip()
    if not name or len(name) > 255:
        raise MaterialError("Provide a material name of at most 255 characters.")
    return name


def _bounded_text(parts, warnings):
    chunks = []
    remaining = MAX_TEXT_CHARS
    for part in parts:
        if not part:
            continue
        if chunks:
            part = "\n" + part
        chunks.append(part[:remaining])
        remaining -= min(len(part), remaining)
        if len(part) > len(chunks[-1]):
            warnings.append("Extracted text truncated at 60000 characters.")
            break
    text = "".join(chunks)
    if not text.strip():
        warnings.append("No extractable text found; OCR is not performed.")
    return text


class _HTMLText(HTMLParser):
    """Drop active/non-content elements and retain only text, never markup."""

    HIDDEN = frozenset(("script", "style", "iframe", "object", "template", "noscript", "svg"))
    BREAKS = frozenset(("p", "div", "br", "li", "tr", "h1", "h2", "h3", "section"))

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.hidden = []
        self.parts = []

    def handle_starttag(self, tag, attrs):
        if tag in self.HIDDEN:
            self.hidden.append(tag)
        elif not self.hidden and tag in self.BREAKS:
            self.parts.append("\n")

    def handle_startendtag(self, tag, attrs):
        if not self.hidden and tag in self.BREAKS:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in self.hidden:
            # Ignore malformed nesting conservatively until this hidden tag ends.
            self.hidden = self.hidden[:self.hidden.index(tag)]
        elif not self.hidden and tag in self.BREAKS:
            self.parts.append("\n")

    def handle_data(self, data):
        if not self.hidden:
            self.parts.append(data)


def _decode(data, warnings):
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        warnings.append("Non-UTF-8 bytes replaced while decoding text.")
        return data.decode("utf-8-sig", errors="replace")


def _zip_preflight(data, warnings):
    """Inspect all entries before a document parser expands any of them."""
    from defusedxml import ElementTree

    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        entries = archive.infolist()
        if len(entries) > MAX_ZIP_ENTRIES:
            raise MaterialError("Document archive has too many entries.")
        expanded = 0
        names = set()
        for entry in entries:
            name = entry.filename
            path = PurePosixPath(name)
            if (name in names or "\\" in name or name.startswith("/")
                    or ".." in path.parts or ":" in name
                    or stat.S_ISLNK(entry.external_attr >> 16) or entry.flag_bits & 1):
                raise MaterialError("Document archive contains unsafe entries.")
            names.add(name)
            expanded += entry.file_size
            if expanded > MAX_ZIP_BYTES:
                raise MaterialError("Document archive exceeds the 32 MiB expanded limit.")
            if entry.file_size > max(1, entry.compress_size) * MAX_ZIP_RATIO:
                raise MaterialError("Document archive compression ratio exceeds the safety limit.")
        external = False
        for entry in entries:
            if entry.filename.endswith((".xml", ".rels")):
                # defusedxml rejects entities; no XML parser gets to resolve a URL.
                root = ElementTree.fromstring(archive.read(entry), forbid_dtd=True)
                if entry.filename.endswith(".rels"):
                    external |= any(e.attrib.get("TargetMode", "").lower() == "external"
                                    for e in root.iter())
        if external:
            warnings.append("External document relationships are not fetched.")


def _extract_parts(data, extension, warnings):
    if extension in (".txt", ".md", ".csv", ".html"):
        text = _decode(data, warnings)
        if extension == ".html":
            parser = _HTMLText()
            parser.feed(text)
            parser.close()
            text = "".join(parser.parts)
        yield text
    elif extension == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise MaterialError("Encrypted PDFs cannot be read; upload an unlocked copy.")
        for page in reader.pages:
            yield page.extract_text() or ""
    else:
        _zip_preflight(data, warnings)
        if extension == ".docx":
            from docx import Document

            document = Document(io.BytesIO(data))
            for paragraph in document.paragraphs:
                yield paragraph.text
            for table in document.tables:
                for row in table.rows:
                    yield "\t".join(cell.text for cell in row.cells)
        elif extension == ".pptx":
            from pptx import Presentation

            for slide in Presentation(io.BytesIO(data)).slides:
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        yield shape.text
                    if shape.has_table:
                        for row in shape.table.rows:
                            yield "\t".join(cell.text for cell in row.cells)


def extract_material(path, original_name=None, role="reference"):
    """Read one regular local file (10 MiB max); never follow a leaf symlink."""
    role = _role(role)
    name = _name(original_name if original_name is not None else path)
    extension = Path(name).suffix.lower()
    if extension not in EXTENSIONS:
        raise MaterialError("Unsupported material type; upload TXT, MD, CSV, HTML, PDF, DOCX, or PPTX.")
    warnings = _warnings(role)
    try:
        flags = os.O_RDONLY | getattr(os, "O_NOFOLLOW", 0) | getattr(os, "O_NONBLOCK", 0)
        descriptor = os.open(path, flags)
        with os.fdopen(descriptor, "rb") as stream:
            info = os.fstat(stream.fileno())
            if not stat.S_ISREG(info.st_mode):
                raise MaterialError("Material must be a regular local file.")
            if info.st_size > MAX_FILE_BYTES:
                raise MaterialError("Material exceeds the 10 MiB file limit.")
            data = stream.read(MAX_FILE_BYTES + 1)
        if len(data) > MAX_FILE_BYTES:
            raise MaterialError("Material exceeds the 10 MiB file limit.")
        text = _bounded_text(_extract_parts(data, extension, warnings), warnings)
    except MaterialError:
        raise
    except ImportError:
        raise MaterialError("The document parser is unavailable; upload plain text instead.") from None
    except Exception:
        raise MaterialError("Cannot read this material; upload a valid unlocked document or plain text.") from None
    return dict(name=name, kind="file", role=role, text=text, warnings=warnings,
                sha256=hashlib.sha256(data).hexdigest(), bytes=len(data))


def text_material(text, name="Instructor notes", role="reference"):
    """Validate instructor notes; oversized direct input is rejected, not lost."""
    role = _role(role)
    name = _name(name)
    if not isinstance(text, str) or not text.strip():
        raise MaterialError("Enter nonempty material text.")
    if len(text) > MAX_TEXT_CHARS:
        raise MaterialError("Material text exceeds the 60000 character limit.")
    try:
        data = text.encode("utf-8")
    except UnicodeError:
        raise MaterialError("Material text contains invalid characters.") from None
    return dict(name=name, kind="text", role=role, text=text, warnings=_warnings(role),
                sha256=hashlib.sha256(data).hexdigest(), bytes=len(data))


def _normalize_url(url):
    if (not isinstance(url, str) or len(url) > 4096 or not url
            or any(c.isspace() or ord(c) < 32 or ord(c) == 127 for c in url)
            or "\\" in url):
        raise MaterialError("Provide a valid public HTTP or HTTPS URL.")
    try:
        parsed = urlsplit(url)
        if (parsed.scheme not in ("http", "https") or not parsed.hostname
                or parsed.username is not None or parsed.password is not None):
            raise ValueError()
        host = parsed.hostname.encode("idna").decode("ascii").lower()
        port = parsed.port if parsed.port is not None else (443 if parsed.scheme == "https" else 80)
        if port != (443 if parsed.scheme == "https" else 80) or "%" in host:
            raise ValueError()
        # Brackets belong to URL authority, not DNS or TLS server names.
        authority = "[" + host + "]" if ":" in host else host
        target = quote(parsed.path or "/", safe="/%:@!$&'()*+,;=-._~")
        query = quote(parsed.query, safe="/?%:@!$&'()*+,;=-._~")
        return urlunsplit((parsed.scheme, authority, target, query, ""))
    except (ValueError, UnicodeError):
        raise MaterialError("Use a public HTTP/HTTPS URL without credentials or a nonstandard port.") from None


def _public_address(host, port):
    """Reject the whole answer if even one result is not public."""
    try:
        results = socket.getaddrinfo(host, port, type=socket.SOCK_STREAM)
        addresses = []
        for family, _, _, _, sockaddr in results:
            if family not in (socket.AF_INET, socket.AF_INET6):
                raise ValueError()
            value = sockaddr[0]
            address = ipaddress.ip_address(value)
            if (not address.is_global or address.is_multicast or address.is_reserved
                    or address.is_loopback or address.is_link_local or address.is_private
                    or address.is_unspecified or "%" in value):
                raise ValueError()
            if address.version == 6 and address.ipv4_mapped is not None:
                mapped = address.ipv4_mapped
                if not mapped.is_global or mapped.is_reserved or mapped.is_multicast:
                    raise ValueError()
            addresses.append(value)
        if not addresses:
            raise ValueError()
        return addresses[0]
    except (OSError, ValueError):
        raise MaterialError("URL must resolve exclusively to public Internet addresses.") from None


class _PinnedHTTPConnection(http.client.HTTPConnection):
    def __init__(self, host, port, address):
        super().__init__(host, port, timeout=CONNECT_TIMEOUT)
        self.address = address

    def connect(self):
        self.sock = socket.create_connection((self.address, self.port), self.timeout)
        self.sock.settimeout(READ_TIMEOUT)


class _PinnedHTTPSConnection(_PinnedHTTPConnection):
    default_port = 443
    def connect(self):
        super().connect()
        try:
            # Use the URL host for SNI and certificate checks, the checked IP for TCP.
            self.sock = ssl.create_default_context().wrap_socket(self.sock, server_hostname=self.host)
        except Exception:
            self.sock.close()
            raise


@contextmanager
def _open_response(url, address):
    parsed = urlsplit(url)
    cls = _PinnedHTTPSConnection if parsed.scheme == "https" else _PinnedHTTPConnection
    connection = cls(parsed.hostname, 443 if parsed.scheme == "https" else 80, address)
    try:
        target = parsed.path + ("?" + parsed.query if parsed.query else "")
        connection.request("GET", target, headers={
            "Accept": "text/plain, text/html, application/pdf",
            "Accept-Encoding": "identity", "User-Agent": "MentorMaterialReader/1.0",
        })
        response = connection.getresponse()
        try:
            yield response
        finally:
            response.close()
    finally:
        connection.close()


def fetch_url_material(url, role="reference"):
    """Read one public page, with DNS-pinned TCP, verified TLS, and no proxies.

    No cookies, credentials, subresources, JavaScript or external relationships
    are used. Callers should offer local upload for authenticated/paywalled pages.
    """
    role = _role(role)
    current = _normalize_url(url)
    warnings = _warnings(role)
    try:
        for redirects in range(MAX_REDIRECTS + 1):
            parsed = urlsplit(current)
            address = _public_address(parsed.hostname, 443 if parsed.scheme == "https" else 80)
            with _open_response(current, address) as response:
                if response.status in (301, 302, 303, 307, 308):
                    if redirects == MAX_REDIRECTS:
                        raise MaterialError("URL exceeds the three redirect limit.")
                    location = response.getheader("Location")
                    if not location:
                        raise MaterialError("URL returned an invalid redirect.")
                    current = _normalize_url(urljoin(current, location))
                    continue
                if response.status in (401, 402, 403, 407):
                    raise MaterialError("Cannot read authenticated or paywalled pages; upload a local copy instead.")
                if response.status != 200:
                    raise MaterialError("Cannot read this URL; upload a local copy instead.")
                encoding = (response.getheader("Content-Encoding") or "identity").lower()
                if encoding != "identity":
                    raise MaterialError("Compressed URL responses are unsupported; upload a local copy instead.")
                media_type = (response.getheader("Content-Type") or "").split(";")[0].strip().lower()
                if media_type not in ("text/plain", "text/html", "application/pdf"):
                    raise MaterialError("URL must return plain text, HTML, or PDF.")
                length = response.getheader("Content-Length")
                if length is not None:
                    try:
                        length = int(length)
                    except ValueError:
                        raise MaterialError("URL returned an invalid response length.") from None
                    if length < 0 or length > MAX_RESPONSE_BYTES:
                        raise MaterialError("URL response exceeds the 2 MiB limit.")
                data = bytearray()
                while True:
                    chunk = response.read(min(65536, MAX_RESPONSE_BYTES + 1 - len(data)))
                    if not chunk:
                        break
                    data.extend(chunk)
                    if len(data) > MAX_RESPONSE_BYTES:
                        raise MaterialError("URL response exceeds the 2 MiB limit.")
                break
        if media_type == "application/pdf":
            # Never derive a filesystem path or extension from an untrusted URL.
            with tempfile.TemporaryDirectory(prefix="mentor-material-") as directory:
                path = Path(directory) / "document.pdf"
                path.write_bytes(data)
                result = extract_material(path, role=role)
            text = result["text"]
            warnings = result["warnings"]
        else:
            extension = ".html" if media_type == "text/html" else ".txt"
            text = _bounded_text(_extract_parts(bytes(data), extension, warnings), warnings)
    except MaterialError:
        raise
    except Exception:
        raise MaterialError("Cannot safely read this URL; upload a local copy instead.") from None
    return dict(name=urlsplit(current).hostname, url=current, kind="url", role=role,
                text=text, warnings=warnings, bytes=len(data))
