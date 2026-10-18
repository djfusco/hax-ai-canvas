"""Supplementary materials tests: synthetic documents, no network or user data."""
from contextlib import contextmanager
import hashlib
import io
from pathlib import Path
import socket
import stat
import sys
import tempfile
import unittest
from unittest.mock import MagicMock, patch
import zipfile

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "python"))
import mentor_materials as materials


def answer(*addresses):
    return [(socket.AF_INET6 if ":" in address else socket.AF_INET,
             socket.SOCK_STREAM, 6, "", (address, 443)) for address in addresses]


class Response:
    def __init__(self, body=b"Public evidence", status=200, headers=None):
        self.status = status
        self.headers = {"Content-Type": "text/plain", **(headers or {})}
        self.body = io.BytesIO(body)
        self.read_calls = 0

    def getheader(self, name):
        return self.headers.get(name)

    def read(self, size):
        self.read_calls += 1
        return self.body.read(size)


class MaterialTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        # Any accidental external connection fails before it reaches the network.
        self.network = patch("socket.create_connection", side_effect=AssertionError("Unexpected network"))
        self.network.start()

    def tearDown(self):
        self.network.stop()
        self.temp.cleanup()

    def file(self, name, content):
        path = self.root / name
        path.write_bytes(content)
        return path

    def archive(self, entries, name="evidence.docx"):
        path = self.root / name
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
            for key, value in entries:
                archive.writestr(key, value)
        return path

    def test_plain_file_types_and_provenance(self):
        data = "Résumé, syllabus\nCafé".encode()
        for extension in ("txt", "md", "csv"):
            with self.subTest(extension=extension):
                result = materials.extract_material(self.file("source." + extension, data),
                                                    original_name="C:\\fakepath\\Notes." + extension,
                                                    role="student")
                self.assertEqual(result["name"], "Notes." + extension)
                self.assertEqual(result["role"], "student")
                self.assertEqual(result["kind"], "file")
                self.assertEqual(result["text"], data.decode())
                self.assertEqual(result["bytes"], len(data))
                self.assertEqual(result["sha256"], hashlib.sha256(data).hexdigest())

    def test_html_is_text_without_active_content(self):
        path = self.file("page.html", b'<p onclick="bad()">Evidence &amp; context</p>'
                         b'<script>secret code</script><style>hidden styles</style>'
                         b'<iframe src="http://localhost">nested</iframe>'
                         b'<img src="http://localhost"><p>More</p>')
        text = materials.extract_material(path)["text"]
        self.assertIn("Evidence & context", text)
        self.assertIn("More", text)
        for bad in ("secret", "hidden", "nested", "localhost", "<p", "onclick"):
            self.assertNotIn(bad, text)

    def test_file_truncates_but_direct_text_rejects_over_limit(self):
        text = "a" * materials.MAX_TEXT_CHARS
        result = materials.extract_material(self.file("long.txt", (text + "tail").encode()))
        self.assertEqual(result["text"], text)
        self.assertTrue(any("truncated" in warning for warning in result["warnings"]))
        self.assertEqual(materials.text_material(text)["text"], text)
        with self.assertRaises(materials.MaterialError):
            materials.text_material(text + "a")

    def test_roles_policy_and_text_validation(self):
        for role in materials.ROLES:
            result = materials.text_material("notes", role=role)
            self.assertEqual(result["role"], role)
            self.assertEqual(bool(result["warnings"]), role == "policy")
        for role in ("admin", "", None, []):
            with self.subTest(role=role), self.assertRaises(materials.MaterialError):
                materials.text_material("notes", role=role)
        for value in ("", " \n", None, b"notes", "\ud800"):
            with self.subTest(value=repr(value)), self.assertRaises(materials.MaterialError):
                materials.text_material(value)

    def test_file_size_unsupported_missing_and_symlink(self):
        path = self.root / "large.txt"
        with path.open("wb") as stream:
            stream.truncate(materials.MAX_FILE_BYTES + 1)
        for value in (path, self.file("program.py", b"print(1)"), self.root / "missing.txt"):
            with self.subTest(value=value), self.assertRaises(materials.MaterialError):
                materials.extract_material(value)
        link = self.root / "link.txt"
        link.symlink_to(self.file("actual.txt", b"not followed"))
        with self.assertRaises(materials.MaterialError):
            materials.extract_material(link)
        with self.assertRaises(materials.MaterialError):
            materials.extract_material(self.root, original_name="directory.txt")

    def test_safe_errors_do_not_expose_parser_exception(self):
        with patch.object(materials, "_extract_parts", side_effect=RuntimeError("private TOKEN body")):
            with self.assertRaises(materials.MaterialError) as caught:
                materials.extract_material(self.file("bad.pdf", b"invalid"))
        self.assertNotIn("TOKEN", str(caught.exception))
        self.assertNotIn(str(self.root), str(caught.exception))

    def test_empty_pdf_and_encrypted_pdf(self):
        from pypdf import PdfWriter
        for encrypted in (False, True):
            writer = PdfWriter()
            writer.add_blank_page(width=72, height=72)
            if encrypted:
                writer.encrypt("fixture-password")
            path = self.root / "empty.pdf"
            with path.open("wb") as stream:
                writer.write(stream)
            if encrypted:
                with self.assertRaisesRegex(materials.MaterialError, "Encrypted"):
                    materials.extract_material(path)
            else:
                result = materials.extract_material(path)
                self.assertEqual(result["text"], "")
                self.assertTrue(any("OCR is not performed" in w for w in result["warnings"]))

    def test_docx_and_external_relationship_does_not_fetch(self):
        from docx import Document
        from docx.opc.constants import RELATIONSHIP_TYPE
        document = Document()
        document.add_paragraph("Document evidence")
        document.add_table(rows=1, cols=1).cell(0, 0).text = "Table evidence"
        document.part.relate_to("http://169.254.169.254/private",
                               RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
        path = self.root / "document.docx"
        document.save(path)
        result = materials.extract_material(path, role="policy")
        self.assertIn("Document evidence", result["text"])
        self.assertIn("Table evidence", result["text"])
        self.assertTrue(any("relationships are not fetched" in w for w in result["warnings"]))

    def test_pptx_text_and_tables(self):
        from pptx import Presentation
        from pptx.util import Inches
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(0, 0, Inches(2), Inches(1)).text = "Slide evidence"
        table = slide.shapes.add_table(1, 1, 0, 0, Inches(2), Inches(1)).table
        table.cell(0, 0).text = "Table evidence"
        path = self.root / "slides.pptx"
        presentation.save(path)
        result = materials.extract_material(path)
        self.assertIn("Slide evidence", result["text"])
        self.assertIn("Table evidence", result["text"])

    def test_document_zip_limits_before_parser(self):
        for extension in ("docx", "pptx"):
            for name in ("../escape.xml", "/absolute.xml", "word\\bad.xml", "C:drive.xml"):
                path = self.archive([(name, "<xml/>")], "bad." + extension)
                with self.subTest(extension=extension, name=name), self.assertRaises(materials.MaterialError):
                    materials.extract_material(path)
        path = self.archive([("word/document.xml", "a" * 100000)])
        with self.assertRaisesRegex(materials.MaterialError, "ratio"):
            materials.extract_material(path)
        path = self.archive([("entry", b"x"), ("entry2", b"x")])
        with patch.object(materials, "MAX_ZIP_ENTRIES", 1):
            with self.assertRaisesRegex(materials.MaterialError, "entries"):
                materials.extract_material(path)
        with patch.object(materials, "MAX_ZIP_BYTES", 1):
            with self.assertRaisesRegex(materials.MaterialError, "expanded"):
                materials.extract_material(path)

    def test_document_rejects_symlinks_and_xml_entities(self):
        entry = zipfile.ZipInfo("link")
        entry.create_system = 3
        entry.external_attr = (stat.S_IFLNK | 0o777) << 16
        path = self.archive([(entry, "../../target")])
        with self.assertRaises(materials.MaterialError):
            materials.extract_material(path)
        path = self.archive([("word/document.xml",
                             '<!DOCTYPE x [<!ENTITY a SYSTEM "http://localhost">]><x>&a;</x>')])
        with self.assertRaises(materials.MaterialError):
            materials.extract_material(path)

    def test_external_dtd_rejected(self):
        path = self.archive([("word/document.xml",
                             '<!DOCTYPE x SYSTEM "http://localhost/private"><x/>')])
        with self.assertRaises(materials.MaterialError):
            materials.extract_material(path)

    def fetch(self, url="https://example.org/page", responses=None, dns=None, role="reference"):
        opened = []
        queue = iter(responses or [Response()])

        @contextmanager
        def transport(target, address):
            opened.append((target, address))
            yield next(queue)

        with patch.object(materials.socket, "getaddrinfo", side_effect=dns,
                          return_value=answer("93.184.216.34")):
            with patch.object(materials, "_open_response", side_effect=transport):
                result = materials.fetch_url_material(url, role=role)
        return result, opened

    def test_public_success_and_html_no_subresources(self):
        response = Response(b"<p>Evidence</p><script>bad()</script><img src='/private'>",
                            headers={"Content-Type": "text/html; charset=utf-8"})
        result, opened = self.fetch(responses=[response], role="policy")
        self.assertEqual(opened, [("https://example.org/page", "93.184.216.34")])
        self.assertEqual(result["url"], "https://example.org/page")
        self.assertEqual(result["kind"], "url")
        self.assertIn("Evidence", result["text"])
        self.assertNotIn("bad()", result["text"])
        self.assertTrue(result["warnings"])

    def test_public_ipv6_and_url_text_truncation(self):
        result, opened = self.fetch(url="https://[2606:4700:4700::1111]/page",
                                   dns=[answer("2606:4700:4700::1111")],
                                   responses=[Response(b"x" * (materials.MAX_TEXT_CHARS + 1))])
        self.assertEqual(opened[0][1], "2606:4700:4700::1111")
        self.assertEqual(len(result["text"]), materials.MAX_TEXT_CHARS)
        self.assertTrue(any("truncated" in w for w in result["warnings"]))

    def test_url_pdf_uses_safe_temporary_name(self):
        from pypdf import PdfWriter
        writer = PdfWriter()
        writer.add_blank_page(width=72, height=72)
        output = io.BytesIO()
        writer.write(output)
        with patch.object(materials, "extract_material", wraps=materials.extract_material) as extractor:
            result, _ = self.fetch(url="https://example.org/../../unsafe.exe",
                                   responses=[Response(output.getvalue(),
                                                       headers={"Content-Type": "application/pdf"})])
        self.assertEqual(Path(extractor.call_args.args[0]).name, "document.pdf")
        self.assertFalse(Path(extractor.call_args.args[0]).exists())
        self.assertEqual(result["text"], "")
        self.assertTrue(any("OCR" in w for w in result["warnings"]))

    def test_all_nonpublic_ipv4_ipv6_and_mixed_dns_rejected(self):
        blocked = ("127.0.0.1", "10.0.0.1", "172.16.0.1", "192.168.1.1",
                   "169.254.169.254", "0.0.0.0", "224.0.0.1", "240.0.0.1",
                   "::1", "fc00::1", "fe80::1", "ff02::1", "2001:db8::1",
                   "::ffff:127.0.0.1")
        for address in blocked:
            with self.subTest(address=address):
                with patch.object(materials.socket, "getaddrinfo", return_value=answer(address)):
                    with patch.object(materials, "_open_response") as transport:
                        with self.assertRaises(materials.MaterialError):
                            materials.fetch_url_material("https://example.org")
                        transport.assert_not_called()
        with self.assertRaises(materials.MaterialError):
            self.fetch(dns=[answer("93.184.216.34", "10.0.0.1")])

    def test_redirects_recheck_dns_including_same_host(self):
        redirect = Response(status=302, headers={"Location": "/next"})
        with self.assertRaises(materials.MaterialError):
            self.fetch(responses=[redirect], dns=[answer("93.184.216.34"), answer("127.0.0.1")])
        with self.assertRaises(materials.MaterialError):
            self.fetch(responses=[Response(status=302, headers={"Location": "http://169.254.169.254/"})],
                       dns=[answer("93.184.216.34"), answer("169.254.169.254")])
        result, opened = self.fetch(responses=[redirect, Response()])
        self.assertEqual(len(opened), 2)
        self.assertEqual(result["url"], "https://example.org/next")

    def test_redirect_limit_and_credentials_redirect(self):
        with self.assertRaisesRegex(materials.MaterialError, "redirect"):
            self.fetch(responses=[Response(status=302, headers={"Location": "/again"}) for _ in range(4)])
        with self.assertRaises(materials.MaterialError):
            self.fetch(responses=[Response(status=302, headers={"Location": "https://user:secret@example.org"})])

    def test_response_limit_headers_and_streaming(self):
        huge = Response(headers={"Content-Length": str(materials.MAX_RESPONSE_BYTES + 1)})
        with self.assertRaisesRegex(materials.MaterialError, "2 MiB"):
            self.fetch(responses=[huge])
        self.assertEqual(huge.read_calls, 0)
        with patch.object(materials, "MAX_RESPONSE_BYTES", 10):
            with self.assertRaisesRegex(materials.MaterialError, "2 MiB"):
                self.fetch(responses=[Response(b"x" * 11)])
            result, _ = self.fetch(responses=[Response(b"x" * 10)])
            self.assertEqual(result["bytes"], 10)
        for length in ("-1", "invalid"):
            with self.subTest(length=length), self.assertRaises(materials.MaterialError):
                self.fetch(responses=[Response(headers={"Content-Length": length})])

    def test_unsupported_response_and_auth_safe_errors(self):
        for headers in ({"Content-Type": "application/json"}, {"Content-Encoding": "gzip"},
                        {"Content-Type": ""}):
            with self.subTest(headers=headers), self.assertRaises(materials.MaterialError):
                self.fetch(responses=[Response(b"SECRET BODY", headers=headers)])
        for status in (401, 402, 403, 407):
            with self.subTest(status=status):
                with self.assertRaisesRegex(materials.MaterialError, "upload") as caught:
                    self.fetch(responses=[Response(b"SECRET BODY", status=status)])
                self.assertNotIn("SECRET", str(caught.exception))
        with patch.object(materials.socket, "getaddrinfo", side_effect=socket.gaierror("SECRET")):
            with self.assertRaises(materials.MaterialError) as caught:
                materials.fetch_url_material("https://example.org")
            self.assertNotIn("SECRET", str(caught.exception))

    def test_invalid_urls_never_reach_dns(self):
        urls = ("file:///etc/passwd", "ftp://example.org", "https://user:secret@example.org",
                "http://example.org:8080", "https://example.org:80", "http://example.org:0",
                "https://example.org/\r\nbad",
                "https://example.org\\@localhost", "https://[fe80::1%25en0]", None, "")
        with patch.object(materials.socket, "getaddrinfo") as resolver:
            for url in urls:
                with self.subTest(url=url), self.assertRaises(materials.MaterialError):
                    materials.fetch_url_material(url)
            resolver.assert_not_called()

    def test_role_validation_before_file_or_url_access(self):
        with patch.object(materials.socket, "getaddrinfo") as resolver:
            with self.assertRaises(materials.MaterialError):
                materials.fetch_url_material("https://example.org", role="unknown")
            resolver.assert_not_called()
        with self.assertRaises(materials.MaterialError):
            materials.extract_material(self.root / "unopened.txt", role="unknown")

    def test_tcp_pinned_and_tls_host_preserved(self):
        raw = MagicMock()
        wrapped = MagicMock()
        context = MagicMock()
        context.wrap_socket.return_value = wrapped
        with patch.object(materials.socket, "create_connection", return_value=raw) as connect:
            with patch.object(materials.ssl, "create_default_context", return_value=context):
                connection = materials._PinnedHTTPSConnection("example.org", 443, "93.184.216.34")
                connection.connect()
        connect.assert_called_once_with(("93.184.216.34", 443), materials.CONNECT_TIMEOUT)
        raw.settimeout.assert_called_once_with(materials.READ_TIMEOUT)
        context.wrap_socket.assert_called_once_with(raw, server_hostname="example.org")
        self.assertIs(connection.sock, wrapped)

    def test_tls_failure_closes_pinned_socket(self):
        raw = MagicMock()
        context = MagicMock()
        context.wrap_socket.side_effect = materials.ssl.SSLError("certificate mismatch")
        with patch.object(materials.socket, "create_connection", return_value=raw):
            with patch.object(materials.ssl, "create_default_context", return_value=context):
                connection = materials._PinnedHTTPSConnection("example.org", 443, "93.184.216.34")
                with self.assertRaises(materials.ssl.SSLError):
                    connection.connect()
        raw.close.assert_called_once()

    def test_transport_host_and_no_credentials_or_proxy_headers(self):
        with patch.object(materials, "_PinnedHTTPSConnection") as cls:
            connection = cls.return_value
            connection.getresponse.return_value = MagicMock()
            with materials._open_response("https://example.org/path?q=1", "93.184.216.34"):
                pass
            cls.assert_called_once_with("example.org", 443, "93.184.216.34")
            args, kwargs = connection.request.call_args
            self.assertEqual(args, ("GET", "/path?q=1"))
            self.assertEqual(kwargs["headers"]["Accept-Encoding"], "identity")
            for header in ("Authorization", "Proxy-Authorization", "Cookie"):
                self.assertNotIn(header, kwargs["headers"])
            connection.close.assert_called_once()


if __name__ == "__main__":
    unittest.main()
